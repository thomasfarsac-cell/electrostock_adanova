from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from uuid import uuid4

from odf.opendocument import load as load_ods_document
from odf.table import Table, TableCell, TableRow
from odf.text import P
from openpyxl import load_workbook

try:
    import xlrd  # type: ignore
except Exception:  # pragma: no cover
    xlrd = None

from app.config import DATA_DIR
from app.utils import normalize_lookup

IMPORTS_DIR = DATA_DIR / 'imports'

IMPORT_FIELD_DEFINITIONS = [
    {'key': 'part_number', 'label': 'Référence', 'required': True},
    {'key': 'name', 'label': 'Désignation', 'required': True},
    {'key': 'category_path', 'label': 'Catégorie (chemin)', 'required': False},
    {'key': 'manufacturer_name', 'label': 'Marque / fabricant', 'required': False},
    {'key': 'package_name', 'label': 'Format / conditionnement', 'required': False},
    {'key': 'location_name', 'label': 'Emplacement', 'required': False},
    {'key': 'quantity', 'label': 'Quantité', 'required': False},
    {'key': 'minimum_quantity', 'label': 'Stock mini', 'required': False},
    {'key': 'datasheet_url', 'label': 'Fiche technique / lien', 'required': False},
    {'key': 'keywords', 'label': 'Mots-clés', 'required': False},
    {'key': 'description', 'label': 'Description', 'required': False},
    {'key': 'notes', 'label': 'Notes', 'required': False},
]

FIELD_ALIASES = {
    'part_number': {'reference', 'ref', 'partnumber', 'part_number', 'pn', 'référence'},
    'name': {'designation', 'désignation', 'name', 'nom', 'component'},
    'category_path': {'categorie', 'catégorie', 'category', 'categorypath', 'category_path', 'famille', 'familles'},
    'manufacturer_name': {'fabricant', 'manufacturer', 'maker', 'vendor'},
    'package_name': {'boitier', 'boîtier', 'package', 'case', 'format'},
    'location_name': {'emplacement', 'location', 'bin', 'drawer', 'storage'},
    'quantity': {'quantite', 'quantité', 'qty', 'quantity', 'stock'},
    'minimum_quantity': {'stockmini', 'minimumquantity', 'minimum_quantity', 'min', 'stock min'},
    'datasheet_url': {'datasheet', 'datasheeturl', 'datasheet_url', 'fichetechnique'},
    'keywords': {'keywords', 'motscles', 'mots-clés', 'tags'},
    'description': {'description', 'desc'},
    'notes': {'notes', 'comment', 'commentaires', 'remarks'},
}

TEMPLATE_HEADERS = [
    'Reference', 'Designation', 'Categorie', 'Fabricant', 'Boitier', 'Quantite',
    'Stock mini', 'Emplacement', 'Datasheet', 'Mots-cles', 'Description', 'Notes',
]

TEMPLATE_SAMPLE_ROW = [
    'VIS-M3-12', 'Vis M3 x 12 mm', 'Quincaillerie / Visserie',
    'Generique', 'Sachet 100 pieces', '25', '5', 'Tiroir A2',
    'https://example.com/fiche-technique.pdf', 'vis;m3;atelier',
    'Vis de montage pour projets atelier', 'Exemple de ligne importable',
]


def ensure_imports_dir():
    IMPORTS_DIR.mkdir(parents=True, exist_ok=True)


def _cell_text(cell: TableCell) -> str:
    chunks = []
    for paragraph in cell.getElementsByType(P):
        text = ''.join(getattr(node, 'data', '') for node in paragraph.childNodes)
        chunks.append(text)
    return '\n'.join(chunk for chunk in chunks if chunk).strip()


def _rows_to_dicts(raw_rows: list[list[str]]):
    cleaned_rows = [list(row) for row in raw_rows if any(str(value).strip() for value in row)]
    if not cleaned_rows:
        raise ValueError('Le fichier est vide.')

    headers = [str(value).strip() or f'Colonne {index + 1}' for index, value in enumerate(cleaned_rows[0])]
    data_rows = []
    for row in cleaned_rows[1:]:
        padded = row + [''] * (len(headers) - len(row))
        data_rows.append({headers[index]: str(padded[index]).strip() for index in range(len(headers))})
    if not data_rows:
        raise ValueError('Le fichier ne contient aucune donnée après l’en-tête.')
    return headers, data_rows


def _parse_csv(content: bytes):
    for encoding in ('utf-8-sig', 'utf-8', 'latin-1'):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError('Encodage CSV non reconnu.')

    sample = text[:2048]
    delimiter = ';'
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=';,\t,')
        delimiter = dialect.delimiter
    except csv.Error:
        if text.count(',') > text.count(';'):
            delimiter = ','

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_rows = [[str(value).strip() for value in row] for row in reader]
    return _rows_to_dicts(raw_rows)


def _parse_xlsx(content: bytes):
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    raw_rows = []
    for row in sheet.iter_rows(values_only=True):
        raw_rows.append(['' if value is None else str(value).strip() for value in row])
    return _rows_to_dicts(raw_rows)


def _parse_xls(content: bytes):
    if xlrd is None:
        raise ValueError('Support XLS indisponible : le module xlrd n’est pas installé.')
    workbook = xlrd.open_workbook(file_contents=content)
    sheet = workbook.sheet_by_index(0)
    raw_rows = []
    for row_index in range(sheet.nrows):
        raw_rows.append(['' if value is None else str(value).strip() for value in sheet.row_values(row_index)])
    return _rows_to_dicts(raw_rows)


def _parse_ods(content: bytes):
    document = load_ods_document(io.BytesIO(content))
    tables = document.spreadsheet.getElementsByType(Table)
    if not tables:
        raise ValueError('Aucune feuille trouvée dans le fichier ODS.')
    sheet = tables[0]
    raw_rows: list[list[str]] = []
    for row in sheet.getElementsByType(TableRow):
        current_row: list[str] = []
        for cell in row.getElementsByType(TableCell):
            repeat = int(cell.getAttribute('numbercolumnsrepeated') or 1)
            value = _cell_text(cell)
            for _ in range(repeat):
                current_row.append(value)
        raw_rows.append(current_row)
    return _rows_to_dicts(raw_rows)


def parse_import_file(filename: str, content: bytes):
    suffix = Path(filename or '').suffix.lower()
    if suffix == '.csv':
        headers, rows = _parse_csv(content)
    elif suffix in {'.xlsx', '.xlsm'}:
        headers, rows = _parse_xlsx(content)
    elif suffix == '.xls':
        headers, rows = _parse_xls(content)
    elif suffix == '.ods':
        headers, rows = _parse_ods(content)
    else:
        raise ValueError('Format non supporté. Utilise CSV, XLS, XLSX ou ODS.')

    return {
        'filename': filename,
        'headers': headers,
        'rows': rows,
        'row_count': len(rows),
        'preview_rows': rows[:10],
    }


def save_import_session(payload: dict) -> str:
    ensure_imports_dir()
    token = uuid4().hex
    path = IMPORTS_DIR / f'{token}.json'
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
    return token


def load_import_session(token: str) -> dict:
    ensure_imports_dir()
    path = IMPORTS_DIR / f'{token}.json'
    if not path.exists():
        raise FileNotFoundError('Session d’import introuvable.')
    return json.loads(path.read_text(encoding='utf-8'))


def delete_import_session(token: str):
    path = IMPORTS_DIR / f'{token}.json'
    if path.exists():
        path.unlink(missing_ok=True)


def suggest_mapping(headers: list[str]):
    normalized_headers = {header: normalize_lookup(header).replace('-', '').replace('_', '').replace('/', '') for header in headers}
    mapping = {}
    for field_key, aliases in FIELD_ALIASES.items():
        wanted = {normalize_lookup(alias).replace('-', '').replace('_', '').replace('/', '') for alias in aliases}
        for header, normalized in normalized_headers.items():
            if normalized in wanted:
                mapping[field_key] = header
                break
    return mapping
